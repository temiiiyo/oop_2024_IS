import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import MySQLdb
from MySQLdb import Error
import os
import datetime
from datetime import datetime as dt
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter


class MedicalClinicApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Медицинская Клиника")
        self.root.geometry("500x400")

        self.db_config = {
            'host': 'localhost',
            'user': 'root',
            'password': 'root',
            'database': 'medical_clinic',
            'port': 3306,
            'charset': 'utf8mb4',
            'use_unicode': True,
            'autocommit': True
        }

        self.connection = None
        self.cursor = None

        self.create_widgets()
        self.connect_to_db()

    def connect_to_db(self):
        """Подключение к базе данных MySQL"""
        try:
            self.connection = MySQLdb.connect(**self.db_config)
            self.cursor = self.connection.cursor()
            messagebox.showinfo("Успех", "Подключение к базе данных установлено")
        except Error as e:
            messagebox.showerror("Ошибка подключения", f"Не удалось подключиться к базе данных:\n{str(e)}")
            self.root.destroy()

    def create_widgets(self):
        """Создание элементов интерфейса"""
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main_frame, text="Медицинская Клиника",
                  font=("Arial", 16, "bold")).pack(pady=(0, 20))

        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=5)

        ttk.Button(button_frame, text="Добавить пациента",
                   command=self.show_add_patient_window, width=25).pack(pady=5)

        ttk.Button(button_frame, text="Добавить диагноз",
                   command=self.show_add_diagnosis_window, width=25).pack(pady=5)

        ttk.Button(button_frame, text="Запланировать приём",
                   command=self.show_add_appointment_window, width=25).pack(pady=5)

        ttk.Button(button_frame, text="Анализировать данные",
                   command=self.analyze_data, width=25).pack(pady=20)

        # Статус бар
        self.status_bar = ttk.Label(self.root, text="Готово",
                                    relief=tk.SUNKEN, anchor=tk.W)
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)

    def show_add_patient_window(self):
        """Окно добавления пациента"""
        dialog = tk.Toplevel(self.root)
        dialog.title("Добавление пациента")
        dialog.geometry("400x250")
        dialog.transient(self.root)
        dialog.grab_set()

        self.center_window(dialog, 400, 250)

        ttk.Label(dialog, text="Добавление нового пациента",
                  font=("Arial", 12, "bold")).pack(pady=(15, 10))

        # Форма
        form_frame = ttk.Frame(dialog)
        form_frame.pack(pady=10, padx=20)

        ttk.Label(form_frame, text="ФИО пациента:").grid(row=0, column=0, sticky=tk.W, pady=10)
        full_name_entry = ttk.Entry(form_frame, width=30)
        full_name_entry.grid(row=0, column=1, pady=10, padx=10)
        full_name_entry.focus_set()

        ttk.Label(form_frame, text="Адрес:").grid(row=1, column=0, sticky=tk.W, pady=10)
        address_entry = ttk.Entry(form_frame, width=30)
        address_entry.grid(row=1, column=1, pady=10, padx=10)

        def save_patient():
            full_name = full_name_entry.get().strip()
            address = address_entry.get().strip()

            if not full_name or not address:
                messagebox.showwarning("Внимание", "Заполните все поля!")
                return

            try:
                self.cursor.execute(
                    "INSERT INTO Patients (full_name, address) VALUES (%s, %s)",
                    (full_name, address)
                )
                self.connection.commit()

                messagebox.showinfo("Успех", "Пациент успешно добавлен")
                dialog.destroy()

            except Error as e:
                messagebox.showerror("Ошибка", f"Ошибка при добавлении пациента:\n{str(e)}")

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=20)

        ttk.Button(btn_frame, text="Сохранить", command=save_patient).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="Отмена", command=dialog.destroy).pack(side=tk.LEFT, padx=10)

        dialog.bind('<Return>', lambda e: save_patient())

    def show_add_diagnosis_window(self):
        """Окно добавления диагноза"""
        dialog = tk.Toplevel(self.root)
        dialog.title("Добавление диагноза")
        dialog.geometry("400x300")
        dialog.transient(self.root)
        dialog.grab_set()

        self.center_window(dialog, 400, 300)

        ttk.Label(dialog, text="Добавление нового диагноза",
                  font=("Arial", 12, "bold")).pack(pady=(15, 10))

        form_frame = ttk.Frame(dialog)
        form_frame.pack(pady=10, padx=20)

        ttk.Label(form_frame, text="Название диагноза:").grid(row=0, column=0, sticky=tk.W, pady=10)
        name_entry = ttk.Entry(form_frame, width=30)
        name_entry.grid(row=0, column=1, pady=10, padx=10)
        name_entry.focus_set()

        ttk.Label(form_frame, text="Лечение:").grid(row=1, column=0, sticky=tk.W, pady=10)
        treatment_entry = tk.Text(form_frame, width=30, height=4)
        treatment_entry.grid(row=1, column=1, pady=10, padx=10)

        def save_diagnosis():
            name = name_entry.get().strip()
            treatment = treatment_entry.get("1.0", tk.END).strip()

            if not name or not treatment:
                messagebox.showwarning("Внимание", "Заполните все поля!")
                return

            try:
                self.cursor.execute(
                    "INSERT INTO Diagnoses (name, treatment) VALUES (%s, %s)",
                    (name, treatment)
                )
                self.connection.commit()

                messagebox.showinfo("Успех", "Диагноз успешно добавлен")
                dialog.destroy()

            except Error as e:
                messagebox.showerror("Ошибка", f"Ошибка при добавлении диагноза:\n{str(e)}")

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=20)

        ttk.Button(btn_frame, text="Сохранить", command=save_diagnosis).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="Отмена", command=dialog.destroy).pack(side=tk.LEFT, padx=10)

        dialog.bind('<Return>', lambda e: save_diagnosis())

    def show_add_appointment_window(self):
        """Окно планирования приема"""
        dialog = tk.Toplevel(self.root)
        dialog.title("Планирование приёма")
        dialog.geometry("500x400")
        dialog.transient(self.root)
        dialog.grab_set()

        self.center_window(dialog, 500, 400)

        ttk.Label(dialog, text="Планирование нового приёма",
                  font=("Arial", 12, "bold")).pack(pady=(15, 10))

        form_frame = ttk.Frame(dialog)
        form_frame.pack(pady=10, padx=20)

        ttk.Label(form_frame, text="Врач:").grid(row=0, column=0, sticky=tk.W, pady=8)
        doctor_combo = ttk.Combobox(form_frame, width=28, state="readonly")
        doctor_combo.grid(row=0, column=1, pady=8, padx=10)

        try:
            self.cursor.execute("SELECT doctor_id, full_name, specialization FROM Doctors")
            doctors = self.cursor.fetchall()
            doctor_list = [f"{d[0]}: {d[1]} ({d[2]})" for d in doctors]
            doctor_combo['values'] = doctor_list
            if doctor_list:
                doctor_combo.set(doctor_list[0])
        except Error as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить список врачей:\n{str(e)}")

        ttk.Label(form_frame, text="Пациент:").grid(row=1, column=0, sticky=tk.W, pady=8)
        patient_combo = ttk.Combobox(form_frame, width=28, state="readonly")
        patient_combo.grid(row=1, column=1, pady=8, padx=10)

        try:
            self.cursor.execute("SELECT patient_id, full_name FROM Patients")
            patients = self.cursor.fetchall()
            patient_list = [f"{p[0]}: {p[1]}" for p in patients]
            patient_combo['values'] = patient_list
            if patient_list:
                patient_combo.set(patient_list[0])
        except Error as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить список пациентов:\n{str(e)}")

        ttk.Label(form_frame, text="Дата:").grid(row=2, column=0, sticky=tk.W, pady=8)
        date_frame = ttk.Frame(form_frame)
        date_frame.grid(row=2, column=1, sticky=tk.W, pady=8, padx=10)

        year_spin = ttk.Spinbox(date_frame, from_=2024, to=2030, width=5)
        year_spin.pack(side=tk.LEFT, padx=2)
        year_spin.set(dt.now().year)

        ttk.Label(date_frame, text="-").pack(side=tk.LEFT)

        month_spin = ttk.Spinbox(date_frame, from_=1, to=12, width=3)
        month_spin.pack(side=tk.LEFT, padx=2)
        month_spin.set(dt.now().month)

        ttk.Label(date_frame, text="-").pack(side=tk.LEFT)

        day_spin = ttk.Spinbox(date_frame, from_=1, to=31, width=3)
        day_spin.pack(side=tk.LEFT, padx=2)
        day_spin.set(dt.now().day)

        ttk.Label(form_frame, text="Время:").grid(row=3, column=0, sticky=tk.W, pady=8)
        time_combo = ttk.Combobox(form_frame, width=28, state="readonly")
        time_combo.grid(row=3, column=1, pady=8, padx=10)
        time_combo['values'] = ["09:00", "10:00", "11:00", "12:00", "13:00", "14:00", "15:00", "16:00", "17:00"]
        time_combo.set("10:00")

        ttk.Label(form_frame, text="Диагноз:").grid(row=4, column=0, sticky=tk.W, pady=8)
        diagnosis_combo = ttk.Combobox(form_frame, width=28, state="readonly")
        diagnosis_combo.grid(row=4, column=1, pady=8, padx=10)

        try:
            self.cursor.execute("SELECT diagnosis_id, name FROM Diagnoses")
            diagnoses = self.cursor.fetchall()
            diagnosis_list = [f"{d[0]}: {d[1]}" for d in diagnoses]
            diagnosis_combo['values'] = diagnosis_list
            if diagnosis_list:
                diagnosis_combo.set(diagnosis_list[0])
        except Error as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить список диагнозов:\n{str(e)}")

        def save_appointment():
            doctor_id = int(doctor_combo.get().split(":")[0]) if doctor_combo.get() else None
            patient_id = int(patient_combo.get().split(":")[0]) if patient_combo.get() else None
            diagnosis_id = int(diagnosis_combo.get().split(":")[0]) if diagnosis_combo.get() else None
            try:
                date_str = f"{year_spin.get()}-{month_spin.get():0>2}-{day_spin.get():0>2}"
                date = dt.strptime(date_str, "%Y-%m-%d").date()
            except ValueError:
                messagebox.showwarning("Внимание", "Некорректная дата!")
                return

            time_str = time_combo.get()

            if not all([doctor_id, patient_id, diagnosis_id, time_str]):
                messagebox.showwarning("Внимание", "Заполните все поля!")
                return

            try:
                self.cursor.execute(
                    "INSERT INTO Appointments (doctor_id, patient_id, date, time, diagnosis_id) VALUES (%s, %s, %s, %s, %s)",
                    (doctor_id, patient_id, date_str, time_str + ":00", diagnosis_id)
                )
                self.connection.commit()

                messagebox.showinfo("Успех", "Приём успешно запланирован")
                dialog.destroy()

            except Error as e:
                messagebox.showerror("Ошибка", f"Ошибка при планировании приёма:\n{str(e)}")

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=20)

        ttk.Button(btn_frame, text="Сохранить", command=save_appointment).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="Отмена", command=dialog.destroy).pack(side=tk.LEFT, padx=10)

    def analyze_data(self):
        """Анализ данных и экспорт в Excel"""
        try:
            # 1. Пациенты с многократными визитами
            self.cursor.execute("""
                SELECT p.patient_id, p.full_name, COUNT(*) as visit_count
                FROM Patients p
                JOIN Appointments a ON p.patient_id = a.patient_id
                GROUP BY p.patient_id, p.full_name
                HAVING COUNT(*) > 1
            """)
            patients_with_multiple_visits = self.cursor.fetchall()

            self.cursor.execute("""
                SELECT d.specialization, COUNT(*) as appointment_count
                FROM Doctors d
                JOIN Appointments a ON d.doctor_id = a.doctor_id
                GROUP BY d.specialization
            """)
            workload_by_specialization = self.cursor.fetchall()

            self.export_to_excel(patients_with_multiple_visits, workload_by_specialization)

            messagebox.showinfo("Успех", "Данные успешно проанализированы и экспортированы в output.xlsx")

        except Error as e:
            messagebox.showerror("Ошибка анализа", f"Ошибка при анализе данных:\n{str(e)}")

    def export_to_excel(self, patients_data, workload_data):
        """Экспорт данных в Excel файл"""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Пациенты_визиты"

        title_cell = ws1.cell(row=1, column=1, value="Пациенты с многократными визитами")
        title_cell.font = Font(bold=True, size=14)
        ws1.merge_cells('A1:C1')

        headers1 = ["ID пациента", "ФИО пациента", "Количество визитов"]
        for col, header in enumerate(headers1, 1):
            cell = ws1.cell(row=2, column=col, value=header)
            cell.fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
            cell.font = Font(bold=True)

        for row, patient in enumerate(patients_data, 3):
            ws1.cell(row=row, column=1, value=patient[0])
            ws1.cell(row=row, column=2, value=patient[1])
            ws1.cell(row=row, column=3, value=patient[2])

        for column in ws1.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws1.column_dimensions[column_letter].width = adjusted_width

        ws2 = wb.create_sheet(title="Загрузка_врачей")

        # Заголовок листа
        title_cell2 = ws2.cell(row=1, column=1, value="Относительная загрузка врачей по специальностям")
        title_cell2.font = Font(bold=True, size=14)
        ws2.merge_cells('A1:B1')

        headers2 = ["Специальность", "Количество приемов"]
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=2, column=col, value=header)
            cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            cell.font = Font(bold=True)

        for row, workload in enumerate(workload_data, 3):
            ws2.cell(row=row, column=1, value=workload[0])
            ws2.cell(row=row, column=2, value=workload[1])

        for column in ws2.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws2.column_dimensions[column_letter].width = adjusted_width

        filename = "output.xlsx"
        wb.save(filename)
        self.status_bar.config(text=f"Файл сохранен: {filename}")

    def center_window(self, window, width, height):
        """Центрирование окна относительно главного окна"""
        window.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() // 2) - (width // 2)
        y = self.root.winfo_y() + (self.root.winfo_height() // 2) - (height // 2)
        window.geometry(f'{width}x{height}+{x}+{y}')

    def __del__(self):
        """Закрытие соединения при завершении"""
        try:
            if hasattr(self, 'cursor') and self.cursor:
                self.cursor.close()
            if hasattr(self, 'connection') and self.connection:
                self.connection.close()
        except:
            pass
def main():
    root = tk.Tk()
    app = MedicalClinicApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()