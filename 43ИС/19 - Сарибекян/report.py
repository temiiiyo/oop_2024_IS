from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


class Report:
    @staticmethod
    def generate_report_from_db(db, filename="unprocessed_requests_db.xlsx"):
        """
        Генерирует отчет из базы данных
        """
        # Получаем данные из базы
        requests_data = db.get_unprocessed_requests()

        if not requests_data:
            print("В базе данных нет неотработанных заявок")
            return

        # Создаем Excel файл
        wb = Workbook()
        ws = wb.active
        ws.title = "Неотработанные заявки"

        # Заголовок
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = "ОТЧЕТ О НЕОТРАБОТАННЫХ ЗАЯВКАХ ЖКХ"
        title_cell.font = Font(size=14, bold=True)
        title_cell.alignment = Alignment(horizontal='center')

        ws.append([])  # Пустая строка

        # Заголовки таблицы
        headers = ["ID заявки", "Заявитель", "Описание", "Статус", "Исполнитель", "Ремонт"]
        ws.append(headers)

        # Стили для заголовков
        for cell in ws[3]:  # Заголовки в строке 3
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Данные
        for row in requests_data:
            ws.append(row)

        # Автоподбор ширины столбцов
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Сохраняем файл
        wb.save(filename)
        print(f"✓ Отчет из базы данных сохранен в: {filename}")
        print(f"✓ Количество записей: {len(requests_data)}")
        return filename

    @staticmethod
    def generate_report_from_objects(requests, filename="unprocessed_requests_objects.xlsx"):
        """
        Генерирует отчет из списка объектов Request
        """
        if not requests:
            print("Нет неотработанных заявок для отчета")
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Неотработанные заявки"

        # Заголовки
        ws.append(["ID", "Заявитель", "Описание", "Статус", "Исполнитель", "Ремонт"])

        # Данные
        for i, request in enumerate(requests, 1):
            ws.append([
                getattr(request, 'id', i),  # Используем ID объекта или порядковый номер
                request.applicant.name,
                request.description,
                request.status,
                request.executor or "Не назначен",
                request.repair or "Не указан"
            ])

        # Автоподбор ширины
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filename)
        print(f"✓ Отчет из объектов сохранен в: {filename}")
        return filename